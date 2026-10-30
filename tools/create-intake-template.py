"""Generate a blank workbook; never embed client documents or real IDs."""
import json,subprocess
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.datavalidation import DataValidation
root=Path(__file__).resolve().parents[1]
model=json.loads(subprocess.check_output(['node','--input-type=module','-e',"import {TABLES,BASIC} from './extension/intake-model.js';console.log(JSON.stringify({TABLES,BASIC}));"],cwd=root))
w=Workbook();w.remove(w.active)
for name,headers in model['TABLES'].items():
 s=w.create_sheet(name);s.append(headers);s.freeze_panes='A2';s.auto_filter.ref=f'A1:{s.cell(1,len(headers)).column_letter}301'
 for c in s[1]:c.font=Font(color='FFFFFF',bold=True);c.fill=PatternFill('solid',fgColor='326C50')
 for i,h in enumerate(headers,1):s.column_dimensions[s.cell(1,i).column_letter].width=24 if h!='내용' else 55
 for row in s.iter_rows(min_row=2,max_row=301,max_col=len(headers)):
  for c in row:c.number_format='@';c.alignment=Alignment(vertical='top',wrap_text=True)
 if name=='기본정보':
  for i,k in enumerate(model['BASIC'],2):s.cell(i,1,k)
 if name=='성분':
  for col,choices in [('C','주,부,혼합,부혼합'),('G','제품,상위성분')]:
   d=DataValidation(type='list',formula1='"'+choices+'"',allow_blank=True);d.errorTitle='선택값 확인';d.error='목록의 값을 선택해주세요.';d.showErrorMessage=True;s.add_data_validation(d);d.add(f'{col}2:{col}301')
w.save(root/'docs/intake-template.xlsx')
import base64
(root/'extension/intake-assets.js').write_text('export const templateBase64='+json.dumps(base64.b64encode((root/'docs/intake-template.xlsx').read_bytes()).decode())+';\nexport const guide='+json.dumps((root/'docs/intake-guide.txt').read_text(),ensure_ascii=False)+';\n')
